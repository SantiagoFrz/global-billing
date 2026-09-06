from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from django.db.models import Sum
from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from rest_framework.decorators import api_view
from rest_framework.response import Response

from apps.billing.models import Installment, ScheduledCharge
from apps.clients.models import Client
from apps.contracts.models import Contract
from apps.expenses.models import Expense, ExpenseAllocation
from apps.payments.models import PaymentAllocation
from apps.provisions.models import ProvisionPlan

from .services import financial_breakdown


def _bucket(days: int) -> str:
    if days <= 0:
        return "Current"
    if days <= 7:
        return "1-7 días"
    if days <= 15:
        return "8-15 días"
    if days <= 30:
        return "16-30 días"
    if days <= 60:
        return "31-60 días"
    return "60+ días"


@api_view(["GET"])
def aging_report(request):
    today = timezone.localdate()
    rows = []
    totals: dict[str, int] = {}
    for item in Installment.objects.filter(is_void=False).select_related(
        "charge__contract__client", "charge__contract"
    ):
        if not item.balance:
            continue
        days = (today - item.due_date).days
        bucket = _bucket(days)
        totals[bucket] = totals.get(bucket, 0) + item.balance
        rows.append(
            {
                "id": str(item.pk),
                "client": item.charge.contract.client.name,
                "amount": item.balance,
                "dueDate": item.due_date,
                "daysOverdue": max(0, days),
                "bucket": bucket,
                "contract": item.charge.contract.name,
            }
        )
    return Response({"asOf": today, "totals": totals, "rows": rows})


@api_view(["GET"])
def projection(request):
    days = min(90, max(30, int(request.query_params.get("days", 30))))
    today = timezone.localdate()
    until = today + timedelta(days=days)
    installments = Installment.objects.filter(is_void=False, due_date__range=(today, until))
    expected_income = sum(item.balance for item in installments)
    expenses = (
        Expense.objects.filter(is_void=False, paid_at__isnull=True, due_date__range=(today, until)).aggregate(
            total=Sum("amount")
        )["total"]
        or 0
    )
    provisions = sum(
        plan.suggested_contribution(today)
        for plan in ProvisionPlan.objects.filter(status=ProvisionPlan.Status.ACTIVE, due_date__lte=until)
    )
    contracts = list(
        Contract.objects.filter(status=Contract.Status.ACTIVE, end_date__range=(today, until)).values(
            "id", "name", "end_date", "client__name"
        )
    )
    current = financial_breakdown(on_date=today)
    return Response(
        {
            "days": days,
            "from": today,
            "until": until,
            "realAvailable": current["available"],
            "expectedIncome": expected_income,
            "futureExpenses": expenses,
            "suggestedProvisions": provisions,
            "projectedAvailable": current["available"] + expected_income - expenses - provisions,
            "contractsExpiring": contracts,
        }
    )


@api_view(["GET"])
def calendar_events(request):
    today = timezone.localdate()
    start = date.fromisoformat(request.query_params.get("from", today.isoformat()))
    end = date.fromisoformat(request.query_params.get("to", (today + timedelta(days=90)).isoformat()))
    events = []
    for item in Installment.objects.filter(is_void=False, due_date__range=(start, end)).select_related(
        "charge__contract__client"
    ):
        events.append(
            {
                "id": str(item.pk),
                "type": "charge",
                "date": item.due_date,
                "title": f"Cobro · {item.charge.contract.client.name}",
                "amount": item.balance,
                "status": item.calculated_status(),
                "href": f"/billing/{item.pk}",
            }
        )
    for item in Expense.objects.filter(is_void=False, due_date__range=(start, end)):
        events.append(
            {
                "id": str(item.pk),
                "type": "expense",
                "date": item.due_date,
                "title": item.concept,
                "amount": item.amount,
                "status": "paid" if item.paid_at else "pending",
                "href": f"/expenses/{item.pk}",
            }
        )
    for item in ProvisionPlan.objects.filter(due_date__range=(start, end)):
        events.append(
            {
                "id": str(item.pk),
                "type": "provision",
                "date": item.due_date,
                "title": item.name,
                "amount": item.remaining_amount,
                "status": item.status,
                "href": f"/provisions/{item.pk}",
            }
        )
    for item in Contract.objects.filter(end_date__range=(start, end)):
        events.append(
            {
                "id": str(item.pk),
                "type": "contract",
                "date": item.end_date,
                "title": f"Vence · {item.name}",
                "amount": None,
                "status": item.status,
                "href": f"/contracts/{item.pk}",
            }
        )
    return Response(sorted(events, key=lambda event: str(event["date"])))


@api_view(["GET"])
def client_profitability(request):
    rows = []
    for client in Client.objects.filter(status=Client.Status.ACTIVE):
        charges = ScheduledCharge.objects.filter(contract__client=client, is_void=False)
        contracted = charges.aggregate(total=Sum("amount"))["total"] or 0
        waived = charges.aggregate(total=Sum("waived_amount"))["total"] or 0
        collectible = contracted - waived
        received = (
            PaymentAllocation.objects.filter(
                payment__client=client, payment__is_void=False, is_void=False
            ).aggregate(total=Sum("amount"))["total"]
            or 0
        )
        direct = (
            Expense.objects.filter(client=client, is_void=False).aggregate(total=Sum("amount"))["total"] or 0
        )
        shared = (
            ExpenseAllocation.objects.filter(client=client, is_void=False).aggregate(total=Sum("amount"))[
                "total"
            ]
            or 0
        )
        reserved = sum(
            plan.reserved_balance
            for plan in client.provision_plans.filter(status=ProvisionPlan.Status.ACTIVE)
        )
        result = received - direct - shared
        rows.append(
            {
                "id": str(client.pk),
                "client": client.name,
                "contracted": contracted,
                "collectible": collectible,
                "received": received,
                "waivers": waived,
                "directCosts": direct,
                "sharedCosts": shared,
                "reserved": reserved,
                "economicResult": result,
                "marginPercentage": round(result / received * 100, 1) if received else 0,
                "cashAfterReservations": result - reserved,
            }
        )
    return Response(rows)


def _monthly_values(year: int, month: int) -> list[tuple[str, int]]:
    data = financial_breakdown(on_date=date(year, month, 1))
    return [
        ("Saldo bancario", data["bankBalance"]),
        ("Dinero reservado", data["reserved"]),
        ("Dinero comprometido", data["committed"]),
        ("Dinero disponible", data["available"]),
        ("Por cobrar", data["receivableThisMonth"]),
        ("Cartera vencida", data["overdue"]),
        ("Resultado del mes", data["monthResult"]),
    ]


@api_view(["GET"])
def monthly_report_xlsx(request):
    today = timezone.localdate()
    year = int(request.query_params.get("year", today.year))
    month = int(request.query_params.get("month", today.month))
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen mensual"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "GLOBAL BILLING · RESUMEN MENSUAL"
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="522E88")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws["A3"], ws["B3"] = "Métrica", "Valor (COP)"
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20172D")
    for row, (label, value) in enumerate(_monthly_values(year, month), start=4):
        ws.cell(row, 1, label)
        ws.cell(row, 2, value).number_format = '"$"#,##0;[Red]("$"#,##0)'
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws["A12"] = f"Periodo: {month:02d}/{year} · Generado {timezone.localtime():%d/%m/%Y %H:%M}"
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename=f"Global_Billing_{year}_{month:02d}.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@api_view(["GET"])
def monthly_report_pdf(request):
    today = timezone.localdate()
    year = int(request.query_params.get("year", today.year))
    month = int(request.query_params.get("month", today.month))
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    width, height = A4
    pdf.setFillColor(HexColor("#522e88"))
    pdf.rect(0, height - 40 * mm, width, 40 * mm, fill=1, stroke=0)
    pdf.setFillColor(HexColor("#ffffff"))
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(20 * mm, height - 23 * mm, "GLOBAL BILLING")
    pdf.setFont("Helvetica", 9)
    pdf.drawString(20 * mm, height - 31 * mm, f"Resumen mensual · {month:02d}/{year}")
    y = height - 62 * mm
    for label, value in _monthly_values(year, month):
        pdf.setFillColor(HexColor("#37313e"))
        pdf.setFont("Helvetica", 10)
        pdf.drawString(22 * mm, y, label)
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawRightString(width - 22 * mm, y, "$" + f"{value:,}".replace(",", "."))
        y -= 12 * mm
    pdf.setFillColor(HexColor("#77717f"))
    pdf.setFont("Helvetica", 8)
    pdf.drawString(20 * mm, 18 * mm, "Cifras administrativas internas · COP · America/Bogota")
    pdf.showPage()
    pdf.save()
    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename=f"Global_Billing_{year}_{month:02d}.pdf",
        content_type="application/pdf",
    )
