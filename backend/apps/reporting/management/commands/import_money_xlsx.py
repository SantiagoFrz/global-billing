from __future__ import annotations

import hashlib
import json
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.reporting.models import ImportBatch, LegacyRecord, ReconciliationIssue


class Command(BaseCommand):
    help = "Analiza o importa Money.xlsx preservando referencias de origen y ambigüedades."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--dry-run", action="store_true")

    def handle(self, *args, **options):
        path = Path(options["path"]).resolve()
        if not path.exists():
            raise CommandError(f"No existe el archivo: {path}")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        wb = load_workbook(path, data_only=True, read_only=False)
        sheet = wb[wb.sheetnames[0]]
        report = self.analyze(sheet)
        report.update(
            {"file": path.name, "sha256": digest, "sheet": sheet.title, "dry_run": bool(options["dry_run"])}
        )
        if options["dry_run"]:
            self.stdout.write(json.dumps(report, ensure_ascii=False, indent=2))
            return
        with transaction.atomic():
            batch, created = ImportBatch.objects.get_or_create(
                sha256=digest,
                status=ImportBatch.Status.IMPORTED,
                defaults={"filename": path.name, "sheet_name": sheet.title, "summary": report},
            )
            if not created:
                self.stdout.write(
                    json.dumps(
                        {**report, "idempotent": True, "batch_id": str(batch.pk)},
                        ensure_ascii=False,
                        indent=2,
                    )
                )
                return
            for record in report["records"]:
                LegacyRecord.objects.create(
                    batch=batch,
                    record_type=record["type"],
                    legacy_period_label=record.get("period", ""),
                    concept=record.get("concept", ""),
                    amount=record.get("amount"),
                    original_value=record,
                    source_sheet=sheet.title,
                    source_row=record["row"],
                    source_cell=record["cell"],
                    requires_review=record.get("requires_review", False),
                    review_reason=record.get("review_reason", ""),
                )
            for issue in report["issues"]:
                ReconciliationIssue.objects.create(batch=batch, **issue)
        self.stdout.write(json.dumps({**report, "batch_id": str(batch.pk)}, ensure_ascii=False, indent=2))

    def analyze(self, sheet) -> dict:
        records: list[dict] = []
        issues: list[dict] = []
        periods: list[str] = []
        for row in range(3, sheet.max_row + 1):
            period = str(sheet.cell(row, 1).value or sheet.cell(row, 7).value or "").strip()
            if not period:
                continue
            periods.append(period)
            ambiguous = not any(ch.isdigit() for ch in period)
            for amount_col, concept_col in ((2, 3), (4, 5)):
                amount = sheet.cell(row, amount_col).value
                concept = sheet.cell(row, concept_col).value
                if amount not in (None, "") or concept not in (None, ""):
                    records.append(self.record("income", period, concept, amount, row, amount_col, ambiguous))
            expense_sum = 0
            for concept_col in range(8, 26, 2):
                amount_col = concept_col + 1
                concept = sheet.cell(row, concept_col).value
                amount = sheet.cell(row, amount_col).value
                if amount not in (None, "") or concept not in (None, ""):
                    value = int(amount or 0)
                    expense_sum += value
                    records.append(self.record("expense", period, concept, value, row, amount_col, ambiguous))
            written_total = sheet.cell(row, 26).value
            written_surplus = sheet.cell(row, 27).value
            income_sum = sum(int(sheet.cell(row, col).value or 0) for col in (2, 4))
            recalculated_surplus = income_sum - expense_sum
            if written_total not in (None, "") and int(written_total) != income_sum:
                issues.append(self.issue("income_total_mismatch", written_total, income_sum, row, "Z"))
            if written_surplus not in (None, "") and int(written_surplus) != recalculated_surplus:
                issues.append(
                    self.issue("surplus_mismatch", written_surplus, recalculated_surplus, row, "AA")
                )
            distribution_sum = 0
            for col, label in zip(range(29, 33), ("Chanty", "Johan", "Global", "Cubillo"), strict=True):
                value = sheet.cell(row, col).value
                if value not in (None, ""):
                    distribution_sum += int(value)
                    records.append(
                        {
                            "type": "legacy_distribution",
                            "period": period,
                            "concept": label,
                            "amount": int(value),
                            "row": row,
                            "cell": f"{get_column_letter(col)}{row}",
                            "requires_review": True,
                            "review_reason": "La parte histórica no se mapea automáticamente a usuarios actuales.",
                        }
                    )
            if (
                written_surplus not in (None, "")
                and distribution_sum
                and distribution_sum != int(written_surplus)
            ):
                issues.append(
                    self.issue("distribution_mismatch", written_surplus, distribution_sum, row, "AC:AF")
                )
            note = sheet.cell(row, 33).value
            if note not in (None, ""):
                records.append(
                    {
                        "type": "legacy_note",
                        "period": period,
                        "concept": "Fondo",
                        "amount": None,
                        "row": row,
                        "cell": f"AG{row}",
                        "original_note": str(note),
                        "requires_review": True,
                        "review_reason": "La nota no representa por sí sola un movimiento financiero.",
                    }
                )
                digits = "".join(ch for ch in str(note) if ch.isdigit())
                if digits and written_surplus not in (None, "") and int(digits) != int(written_surplus):
                    issues.append(
                        self.issue("legacy_note_mismatch", int(digits), int(written_surplus), row, "AG")
                    )
        return {
            "rows_detected": len(set(r["row"] for r in records)),
            "periods": periods,
            "income_count": sum(r["type"] == "income" for r in records),
            "expense_count": sum(r["type"] == "expense" for r in records),
            "distribution_count": sum(r["type"] == "legacy_distribution" for r in records),
            "unknown_concepts": sorted(
                {
                    str(r.get("concept"))
                    for r in records
                    if r["type"] in {"income", "expense"} and r.get("concept")
                }
            ),
            "records": records,
            "issues": issues,
            "totals": {
                "income": sum(r.get("amount") or 0 for r in records if r["type"] == "income"),
                "expense": sum(r.get("amount") or 0 for r in records if r["type"] == "expense"),
                "legacy_distribution": sum(
                    r.get("amount") or 0 for r in records if r["type"] == "legacy_distribution"
                ),
            },
        }

    @staticmethod
    def record(kind, period, concept, amount, row, col, ambiguous):
        return {
            "type": kind,
            "period": period,
            "concept": str(concept or ""),
            "amount": int(amount or 0),
            "row": row,
            "cell": f"{get_column_letter(col)}{row}",
            "requires_review": ambiguous,
            "review_reason": "El periodo no contiene un año inequívoco." if ambiguous else "",
        }

    @staticmethod
    def issue(code, legacy, recalculated, row, col):
        legacy = int(legacy)
        recalculated = int(recalculated)
        return {
            "code": code,
            "severity": "warning",
            "legacy_value": legacy,
            "recalculated_value": recalculated,
            "difference": legacy - recalculated,
            "message": "El valor legado no coincide con el valor recalculado. Ambos se conservan.",
            "source_reference": f"Hoja 1!{col}{row}",
        }
